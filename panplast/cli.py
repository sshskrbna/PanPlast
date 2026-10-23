"""PanPlast: inverted-repeat boundaries, junction genes and core / soft-core genes
of chloroplast genomes, with publication-ready figures, tables and an HTML report.

Example:
    panplast --genbank geseq_results/ --fasta genomes/ --meta metadata.tsv --out results/
"""
from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
import time
from pathlib import Path

import pandas as pd
from Bio import SeqIO

from . import __version__, analysis, plots, report
from .annotate import annotate_fastas
from .genes import accession_of, find_files, genes_from_record, load_synonyms, read_genbank, regions_from_record
from .text import t
from .ir import REGIONS, detect_quadripartite, from_annotation
from .reconcile import reconcile_trna_names

GB_EXT = (".gb", ".gbk", ".genbank", ".gbff")
FA_EXT = (".fa", ".fasta", ".fas", ".fna")


def parse_args(argv=None):
    p = argparse.ArgumentParser(prog="panplast", description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    io = p.add_argument_group("input / output")
    io.add_argument("--genbank", nargs="*", default=[], metavar="PATH",
                    help="annotated GenBank files or folders (GeSeq, PGA, NCBI, ...)")
    io.add_argument("--fasta", nargs="*", default=[], metavar="PATH",
                    help="FASTA files or folders; genomes without a GenBank are annotated")
    io.add_argument("--meta", metavar="TSV",
                    help="table with columns accession (or file), species, optional group/subtribe")
    io.add_argument("--out", default="panplast_results", metavar="DIR", help="output folder [panplast_results]")

    an = p.add_argument_group("annotation of FASTA")
    an.add_argument("--ref", nargs="*", default=[], metavar="GB",
                    help="reference GenBank files for annotating FASTA (default: first 3 of --genbank)")
    an.add_argument("--pga", metavar="PGA.pl", help="annotate FASTA with PGA instead of the built-in blastn transfer")

    ge = p.add_argument_group("genes")
    ge.add_argument("--core", type=float, default=1.0, help="core threshold, fraction of genomes [1.0]")
    ge.add_argument("--soft-core", type=float, default=0.95, help="soft-core threshold [0.95]")
    ge.add_argument("--shell", type=float, default=0.15, help="shell threshold [0.15]")
    ge.add_argument("--synonyms", metavar="TSV", help="extra gene-name synonyms: old_name<TAB>new_name")
    ge.add_argument("--exclude", default="", metavar="LIST", help="comma-separated genes to ignore, e.g. ycf68,ycf15")
    ge.add_argument("--include-pseudo", action="store_true", help="count genes marked /pseudo")
    ge.add_argument("--min-cds-frac", type=float, default=0.5,
                    help="a CDS shorter than this fraction of its typical length is a remnant (partial) [0.5]")

    bo = p.add_argument_group("boundaries and figures")
    bo.add_argument("--ir-source", choices=["blast", "genbank"], default="blast",
                    help="IR boundaries from self-BLAST (default) or from the GenBank annotation")
    bo.add_argument("--min-ir", type=int, default=5000,
                    help="shortest inverted repeat accepted as the IR, bp [5000]; lower it for plastomes with reduced IRs")
    bo.add_argument("--window", type=int, default=1200, help="bp shown on each side of a junction [1200]")
    bo.add_argument("--maps", type=int, default=-1, help="number of circular maps to draw (-1 = all, 0 = none)")
    p.add_argument("--version", action="version", version=f"panplast {__version__}")
    return p.parse_args(argv)


def _label_from_name(stem: str, acc: str | None) -> str:
    s = stem
    for junk in ("_GenBank", "_genbank", "_annotated"):
        s = s.replace(junk, "")
    if acc and acc in s:
        s = s.split(acc, 1)[1]
    return s.strip("_ .-").replace("_", " ")


def _organism_label(org: str, gid: str, stem: str) -> str:
    """The /organism of a record, or the file name when it is missing, a bare '.', an accession or the id."""
    org = org.strip()
    if org.strip(" .") and accession_of(org) is None and org != gid:
        return org
    return _label_from_name(stem, gid)


def _unique_id(gid: str, taken) -> str:
    """gid, or gid_2, gid_3, ... when another genome already uses it."""
    if gid not in taken:
        return gid
    k = 2
    while f"{gid}_{k}" in taken:
        k += 1
    return f"{gid}_{k}"


def load_meta(path):
    if not path:
        return {}
    meta = {}
    with open(path) as fh:
        for row in csv.DictReader(fh, delimiter="\t"):
            row = {k.strip().lower(): (v or "").strip() for k, v in row.items() if k}
            key = row.get("accession") or accession_of(row.get("file", "")) or row.get("file")
            if key:
                meta[key] = dict(species=row.get("species", ""), group=row.get("group") or row.get("subtribe") or "")
    return meta


def main(argv=None):
    a = parse_args(argv)
    t0 = time.time()
    out = Path(a.out)
    (out / "tables").mkdir(parents=True, exist_ok=True)
    (out / "figures").mkdir(parents=True, exist_ok=True)
    log_lines = []

    def log(msg):
        print(msg, flush=True)
        log_lines.append(msg)

    if shutil.which("blastn") is None:
        sys.exit(t("no_blast"))
    syn = load_synonyms(a.synonyms)
    exclude = {x.strip() for x in a.exclude.split(",") if x.strip()}
    meta = load_meta(a.meta)

    # ---------------------------------------------------------------- inputs
    records, source, labels = {}, {}, {}
    gb_files = find_files(a.genbank, GB_EXT)
    for f in gb_files:
        rec = read_genbank(f)
        gid = accession_of(f.name) or accession_of(rec.id) or rec.id
        uid = _unique_id(gid, records)
        if uid != gid:
            log(t("log_dup_id", f=f.name, g=gid, n=uid))
        gid = uid
        records[gid] = rec
        source[gid] = "GenBank: " + f.name
        labels[gid] = _organism_label(rec.annotations.get("organism", ""), gid, f.stem)
    log(t("log_gb", n=len(gb_files)))

    fa_files = find_files(a.fasta, FA_EXT)
    to_annotate = []
    for f in fa_files:
        rec = SeqIO.read(f, "fasta")
        gid = accession_of(rec.id) or accession_of(f.name) or rec.id
        if not labels.get(gid):
            labels[gid] = _label_from_name(rec.description, gid) or _label_from_name(f.stem, gid)
        if gid not in records:
            to_annotate.append(f)
    log(t("log_fa", n=len(fa_files), k=len(to_annotate)))
    if to_annotate:
        refs = find_files(a.ref, GB_EXT) or gb_files[:3]
        if not refs:
            sys.exit(t("no_ref"))
        how = "PGA" if a.pga else t("blast_transfer")
        log(t("log_annot", n=len(to_annotate), how=how, r=len(refs)))
        for p in annotate_fastas(to_annotate, refs, out / "annotated_genbank", a.synonyms, a.pga, log):
            rec = read_genbank(p)
            gid = accession_of(p.name) or rec.id
            uid = _unique_id(gid, records)
            if uid != gid:
                log(t("log_dup_id", f=p.name, g=gid, n=uid))
                labels.setdefault(uid, labels.get(gid, ""))
            gid = uid
            records[gid] = rec
            source[gid] = f"PanPlast {how}" if not a.pga else "PGA"
    if not records:
        sys.exit(t("no_input"))

    for gid, m in meta.items():
        if gid in records and m.get("species"):
            labels[gid] = m["species"]
    groups = {gid: meta.get(gid, {}).get("group", "") for gid in records}
    order = sorted(records, key=lambda g: (groups.get(g, ""), labels.get(g, g)))

    # ---------------------------------------------------------------- per genome
    loci_by, seqs, quads, quads_gb = {}, {}, {}, {}
    for gid in order:
        rec = records[gid]
        seq = str(rec.seq).upper()
        seqs[gid] = seq
        loci_by[gid] = genes_from_record(rec, syn, exclude, a.include_pseudo, a.min_cds_frac)
        q_gb = from_annotation(regions_from_record(rec), len(seq))
        if q_gb and abs(sum(q_gb.region_len(r) for r in REGIONS) - len(seq)) > 10:
            q_gb = None   # broken region annotation
        q = detect_quadripartite(seq, a.min_ir) if a.ir_source == "blast" else q_gb
        if q is None and a.ir_source == "blast" and q_gb:
            q = q_gb
        quads[gid], quads_gb[gid] = q, q_gb
        log(f"  {gid:<12} {labels[gid][:38]:<38} {len(seq):>7} bp  IR={q.region_len('IRb') if q else '—'}")

    renames = reconcile_trna_names(loci_by, seqs, log=log)

    # ---------------------------------------------------------------- tables
    summary = pd.DataFrame([analysis.genome_summary(g, labels[g], seqs[g], loci_by[g], quads[g], quads_gb[g], source[g])
                            for g in order])
    summary.insert(2, "group", [groups.get(g, "") for g in order])
    jrows = []
    for g in order:
        if quads[g]:
            for r in analysis.junction_genes(loci_by[g], quads[g], window=max(3000, a.window)):
                jrows.append(dict(genome=g, species=labels[g], **r))
    junctions = pd.DataFrame(jrows, columns=["genome", "species", "junction", "gene", "fragment", "strand", "rel_start",
                                             "rel_end", "relation", "distance_bp", "bp_left", "bp_right",
                                             "gene_start", "gene_end"])
    matrix = analysis.presence_matrix(loci_by)[order]
    cats = analysis.categorize(matrix, a.core, a.soft_core, a.shell)
    matrix = matrix.loc[cats.gene]
    rare_core, rare_pan = analysis.rarefaction(matrix)
    genes_tab = analysis.gene_lengths(loci_by)
    coords = []
    for g in order:
        q = quads[g]
        if not q:
            continue
        row = dict(genome=g, species=labels[g])
        for r in REGIONS:
            s, e = q.regions[r]
            row[f"{r}_start"], row[f"{r}_end"] = s + 1, e if e > 0 else q.length
        for jn, pos in q.junctions.items():
            row[jn] = pos + 1
        coords.append(row)
    coords = pd.DataFrame(coords)

    qc = []

    def flag(g, level, issue):
        qc.append(dict(genome=g, species=labels[g], level=t(level), issue=issue))

    soft_set = set(cats.gene[cats.percent >= 100 * a.soft_core])
    for g in order:
        q, qg = quads[g], quads_gb[g]
        if not q:
            flag(g, "lvl_error", t("qc_no_ir"))
        if qg is None:
            flag(g, "lvl_note", t("qc_no_gb_regions"))
        elif q and q.source == "blast":
            d = {r: q.region_len(r) - qg.region_len(r) for r in REGIONS}
            if any(abs(v) > 20 for v in d.values()):
                flag(g, "lvl_note", t("qc_ir_diff", d=", ".join(f"{k} {v:+d}" for k, v in d.items() if v)))
        nfrac = summary.set_index("genome").at[g, "N_percent"]
        if nfrac > 1:
            flag(g, "lvl_error", t("qc_n", p=nfrac))
        if q:
            for msg in analysis.orientation_issues(loci_by[g], q):
                flag(g, "lvl_check", msg)
        missing = sorted(gname for gname in soft_set if matrix.at[gname, g] == 0)
        if missing:
            flag(g, "lvl_check", t("qc_missing", n=len(missing), g=", ".join(missing)))
        odd = genes_tab[(genes_tab.genome == g) & genes_tab["flags"].str.contains("nonstandard_tRNA")]
        if len(odd):
            flag(g, "lvl_note", t("qc_trna", g=", ".join(sorted(set(odd.gene)))))
    qc = pd.DataFrame(qc, columns=["genome", "species", "level", "issue"])
    renames = pd.DataFrame(renames, columns=["genome", "start", "old", "new", "identity", "coverage"])

    matrix_out = matrix.copy()
    matrix_out.columns = [f"{labels[g]} ({g})" for g in matrix.columns]
    matrix_out.insert(0, "category", cats.set_index("gene").loc[matrix.index, "category"].astype(str).values)
    tables = {
        "summary": summary, "categories": cats, "matrix": matrix_out.reset_index(names="gene"),
        "junction_genes": junctions, "region_coords": coords, "genes_all": genes_tab, "qc": qc,
        "trna_renamed": renames,
    }
    for k, df in tables.items():
        df.to_csv(out / "tables" / f"{k}.tsv", sep="\t", index=False)
    write_excel(out / "panplast_results.xlsx", tables)
    log(t("log_tables", p=out / "panplast_results.xlsx"))

    # ---------------------------------------------------------------- figures
    figs = {}
    figs["junctions"] = plots.junction_plot(order, labels, quads, jrows, out / "figures" / "ir_junctions", a.window)
    figs["regions"] = plots.region_lengths_plot(order, labels, summary, out / "figures" / "region_lengths")
    figs["pangenome"] = plots.pangenome_plot(cats, matrix, rare_core, rare_pan, out / "figures" / "core_softcore",
                                             (a.core, a.soft_core, a.shell))
    figs["heatmap"] = plots.gene_heatmap(matrix, cats, order, labels, out / "figures" / "gene_matrix")
    figs["heatmap_var"] = plots.gene_heatmap(matrix, cats, order, labels, out / "figures" / "gene_matrix_variable", True)
    nmaps = len(order) if a.maps < 0 else a.maps
    maps = []
    if nmaps:
        (out / "figures" / "maps").mkdir(exist_ok=True)
        for g in order[:nmaps]:
            safe = re.sub(r"[^\w.-]+", "_", f"{g}_{labels[g]}")
            maps.append((g, plots.circular_map(g, labels[g], seqs[g], loci_by[g], quads[g],
                                               out / "figures" / "maps" / safe)))
    log(t("log_figs", p=out / "figures"))

    report.write_report(out / "report.html", summary, cats, junctions, qc, renames, figs, maps, labels,
                        (a.core, a.soft_core, a.shell), out)
    log(t("log_report", p=out / "report.html", s=time.time() - t0))
    (out / "run.log").write_text("\n".join(log_lines) + "\n")


def write_excel(path, tables):
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    names = {"summary": t("sh_summary"), "categories": t("sh_categories"), "matrix": t("sh_matrix"),
             "junction_genes": t("sh_junction"), "region_coords": t("sh_coords"),
             "genes_all": t("sh_genes"), "qc": t("sh_qc"), "trna_renamed": t("sh_renamed")}
    fills = {"core": "D6E6EC", "soft-core": "E3F1F6", "shell": "F7EBD0", "cloud": "ECEDEF"}
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for k, df in tables.items():
            df.to_excel(xw, sheet_name=names[k], index=False)
            ws = xw.sheets[names[k]]
            ws.freeze_panes = "B2" if k == "matrix" else "A2"
            for c in ws[1]:
                c.font = Font(bold=True)
                c.alignment = Alignment(wrap_text=True, vertical="top")
            for i, col in enumerate(df.columns, 1):
                q90 = df[col].astype(str).str.len().quantile(0.9) if len(df) else 8
                q90 = 8 if pd.isna(q90) else int(q90)
                width = min(60, max(8, q90 + 2, len(str(col)) // 2 + 2))
                ws.column_dimensions[get_column_letter(i)].width = width if k != "matrix" or i <= 2 else 5
            if k == "matrix" and len(df):
                rng = f"C2:{get_column_letter(df.shape[1])}{df.shape[0] + 1}"
                ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0, start_color="F1F2F4",
                                                                  mid_type="num", mid_value=1, mid_color="8FBCD4",
                                                                  end_type="num", end_value=2, end_color="2F6F95"))
                ws.row_dimensions[1].height = 140
                for c in ws[1][2:]:
                    c.alignment = Alignment(text_rotation=90)
            if k in ("categories", "matrix"):
                ci = list(df.columns).index("category") + 1
                for row in ws.iter_rows(min_row=2):
                    cat = str(row[ci - 1].value)
                    if cat in fills:
                        row[0].fill = PatternFill("solid", fgColor=fills[cat])
                        row[ci - 1].fill = PatternFill("solid", fgColor=fills[cat])


if __name__ == "__main__":
    main()
